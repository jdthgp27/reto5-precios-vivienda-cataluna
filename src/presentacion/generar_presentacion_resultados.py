# generar_presentacion_resultados.py
# Genera la presentacion PPTX de resultados del proyecto.
#
# Uso:
#   python src/presentacion/generar_presentacion_resultados.py
#
# Salida:
#   reports/presentacion/presentacion_2_resultados.pptx

from pathlib import Path
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

sys.path.append(str(Path(__file__).resolve().parents[2]))
from src.utilidades.config import BASE_DIR, FIGURAS_DIR


SALIDA_PPTX = BASE_DIR / "reports" / "presentacion" / "presentacion_2_resultados.pptx"

# ============================================================
# DISENO (mismo estilo que presentacion 1)
# ============================================================
COLOR_PRIMARIO = RGBColor(0x2E, 0x86, 0xAB)
COLOR_ACENTO = RGBColor(0xF1, 0x8F, 0x01)
COLOR_FONDO = RGBColor(0xF5, 0xF7, 0xFA)
COLOR_TEXTO = RGBColor(0x2C, 0x3E, 0x50)
COLOR_BLANCO = RGBColor(0xFF, 0xFF, 0xFF)
COLOR_ROJO = RGBColor(0xC0, 0x39, 0x2B)
COLOR_VERDE = RGBColor(0x27, 0xAE, 0x60)

FUENTE_TITULO = "Calibri"
FUENTE_TEXTO = "Calibri"


# ============================================================
# UTILIDADES (reutilizadas)
# ============================================================
def nuevo_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def anadir_barra_lateral(slide, prs):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.25), prs.slide_height
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = COLOR_PRIMARIO
    shape.line.fill.background()


def anadir_titulo(slide, texto, color=COLOR_PRIMARIO, tamano=28):
    txBox = slide.shapes.add_textbox(Inches(0.7), Inches(0.4),
                                       Inches(12), Inches(0.9))
    tf = txBox.text_frame
    tf.text = texto
    p = tf.paragraphs[0]
    p.font.name = FUENTE_TITULO
    p.font.size = Pt(tamano)
    p.font.bold = True
    p.font.color.rgb = color


def anadir_subtitulo(slide, texto, tamano=16):
    txBox = slide.shapes.add_textbox(Inches(0.7), Inches(1.3),
                                       Inches(12), Inches(0.6))
    tf = txBox.text_frame
    tf.text = texto
    p = tf.paragraphs[0]
    p.font.name = FUENTE_TEXTO
    p.font.size = Pt(tamano)
    p.font.italic = True
    p.font.color.rgb = COLOR_TEXTO


def anadir_bullets(slide, items, left=0.7, top=2.0, width=12, height=5.5,
                    tamano=14):
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top),
                                       Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True

    for i, item in enumerate(items):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()

        nivel = 0
        if item.startswith("  - "):
            nivel = 1
            item = item[4:]
        elif item.startswith("- "):
            item = item[2:]

        p.text = ("• " if nivel == 0 else "  ◦ ") + item
        p.font.name = FUENTE_TEXTO
        p.font.size = Pt(tamano if nivel == 0 else tamano - 2)
        p.font.color.rgb = COLOR_TEXTO
        p.space_after = Pt(6)
        p.level = nivel


def anadir_imagen(slide, nombre_imagen, left=1.5, top=2.0, height=4.8):
    """Anade una imagen centrada."""
    ruta = FIGURAS_DIR / nombre_imagen
    if not ruta.exists():
        print(f"  AVISO: no existe {ruta}")
        return
    slide.shapes.add_picture(str(ruta), left=Inches(left), top=Inches(top),
                              height=Inches(height))


def anadir_caja_destacada(slide, texto, top=2.0, left=1.0, width=11, height=1.2,
                            color_fondo=None):
    if color_fondo is None:
        color_fondo = COLOR_PRIMARIO
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color_fondo
    shape.line.fill.background()

    tf = shape.text_frame
    tf.word_wrap = True
    tf.text = texto
    p = tf.paragraphs[0]
    p.font.name = FUENTE_TEXTO
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = COLOR_BLANCO
    p.alignment = PP_ALIGN.CENTER


# ============================================================
# SLIDES
# ============================================================
def slide_01_portada(prs):
    slide = nuevo_slide(prs)
    anadir_barra_lateral(slide, prs)

    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.25), 0,
        prs.slide_width - Inches(0.25), Inches(3.5)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = COLOR_ACENTO
    shape.line.fill.background()

    txBox = slide.shapes.add_textbox(Inches(1), Inches(0.8),
                                       Inches(11), Inches(1.2))
    tf = txBox.text_frame
    tf.text = "Resultados e Interpretación"
    p = tf.paragraphs[0]
    p.font.name = FUENTE_TITULO
    p.font.size = Pt(42)
    p.font.bold = True
    p.font.color.rgb = COLOR_BLANCO

    txBox2 = slide.shapes.add_textbox(Inches(1), Inches(2.0),
                                        Inches(11), Inches(0.8))
    tf2 = txBox2.text_frame
    tf2.text = "Predicción de transacciones inmobiliarias · Barcelona 2012-2025"
    p2 = tf2.paragraphs[0]
    p2.font.name = FUENTE_TEXTO
    p2.font.size = Pt(18)
    p2.font.color.rgb = COLOR_BLANCO

    txBox3 = slide.shapes.add_textbox(Inches(1), Inches(4.5),
                                        Inches(11), Inches(2))
    tf3 = txBox3.text_frame
    tf3.text = "Reto 5 · Business Intelligence y Big Data"
    p3 = tf3.paragraphs[0]
    p3.font.size = Pt(18)
    p3.font.bold = True
    p3.font.color.rgb = COLOR_TEXTO

    p4 = tf3.add_paragraph()
    p4.text = "Curso Odisea Data · 2025"
    p4.font.size = Pt(14)
    p4.font.color.rgb = COLOR_TEXTO


def slide_02_indice(prs):
    slide = nuevo_slide(prs)
    anadir_barra_lateral(slide, prs)
    anadir_titulo(slide, "Índice de la presentación")

    items = [
        "1. Resumen ejecutivo (5 hallazgos clave)",
        "2. EDA: Ciclo 2012-2025",
        "3. EDA: Estacionalidad",
        "4. EDA: Ranking de barrios",
        "5. EDA: Distribución por tipologías",
        "6. Análisis SERPAVI vs transacciones",
        "7. Análisis extranjeros por distrito",
        "8. Modelo final: XGBoost",
        "9. SHAP: top features",
        "10. SHAP: caso waterfall",
        "11. Análisis de residuos",
        "12. Limitaciones del modelo",
        "13. Recomendaciones para inversores",
        "14. Recomendaciones para administraciones",
        "15. Cierre",
    ]
    anadir_bullets(slide, items, top=1.8, tamano=13)


def slide_03_resumen_ejecutivo(prs):
    slide = nuevo_slide(prs)
    anadir_barra_lateral(slide, prs)
    anadir_titulo(slide, "1. Resumen ejecutivo")

    anadir_caja_destacada(
        slide,
        "Modelo XGBoost con R² = 0,617 · RMSE = 15,77 · MAPE = 36,7%",
        top=1.5, height=1.0, color_fondo=COLOR_PRIMARIO,
    )

    items = [
        "5 hallazgos clave:",
        "",
        "1. El mercado inmobiliario de Barcelona ha vivido un ciclo completo:",
        "   crisis (2012), COVID (2020) y máximos históricos (2025).",
        "",
        "2. La tendencia histórica del barrio es el factor más predictivo",
        "   (rolling_12m con 23,7% del impacto SHAP).",
        "",
        "3. El precio del alquiler NO predice el volumen de transacciones",
        "   (correlación nula, r = -0,07).",
        "",
        "4. La población extranjera es mayoritariamente inquilina (70,5%),",
        "   especialmente en Ciutat Vella (44,9% estimado).",
        "",
        "5. El modelo funciona bien en valores típicos (15-60 transacciones)",
        "   pero subestima sistemáticamente los picos (100+).",
    ]
    anadir_bullets(slide, items, top=2.7, tamano=12)


def slide_04_ciclo(prs):
    slide = nuevo_slide(prs)
    anadir_barra_lateral(slide, prs)
    anadir_titulo(slide, "2. EDA: Ciclo 2012-2025")

    anadir_imagen(slide, "01_serie_temporal_global.png",
                  left=0.5, top=1.5, height=5.0)

    txBox = slide.shapes.add_textbox(Inches(11.5), Inches(1.8),
                                       Inches(1.7), Inches(4))
    tf = txBox.text_frame
    tf.word_wrap = True
    tf.text = "254.732\ntransacciones\n\n2012-2013:\ncrisis\n\n2020:\nCOVID (-30%)\n\n2025:\nmáximos\nhistóricos"
    for p in tf.paragraphs:
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = COLOR_TEXTO
        p.alignment = PP_ALIGN.CENTER


def slide_05_estacionalidad(prs):
    slide = nuevo_slide(prs)
    anadir_barra_lateral(slide, prs)
    anadir_titulo(slide, "3. EDA: Estacionalidad")

    anadir_imagen(slide, "03_heatmap_anio_mes.png",
                  left=0.5, top=1.5, height=5.0)

    txBox = slide.shapes.add_textbox(Inches(11.5), Inches(1.8),
                                       Inches(1.7), Inches(4))
    tf = txBox.text_frame
    tf.word_wrap = True
    tf.text = "Julio:\nmáximo\n\nAgosto:\nmínimo\n\nPico en julio\ndebido al\nretraso de\nlas escrituras\nnotariales"
    for p in tf.paragraphs:
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = COLOR_TEXTO
        p.alignment = PP_ALIGN.CENTER


def slide_06_ranking(prs):
    slide = nuevo_slide(prs)
    anadir_barra_lateral(slide, prs)
    anadir_titulo(slide, "4. EDA: Ranking de barrios")

    anadir_imagen(slide, "04_ranking_barrios.png",
                  left=0.3, top=1.5, height=5.2)

    txBox = slide.shapes.add_textbox(Inches(11.5), Inches(1.8),
                                       Inches(1.7), Inches(4))
    tf = txBox.text_frame
    tf.word_wrap = True
    tf.text = "Top 3:\n\n1. Sant\nGervasi -\nGalvany\n(11.086)\n\n2. Nova\nEsq.\nEixample\n(10.550)\n\n3. Dreta\nEixample\n(9.900)"
    for p in tf.paragraphs:
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = COLOR_TEXTO
        p.alignment = PP_ALIGN.CENTER


def slide_07_tipologias(prs):
    slide = nuevo_slide(prs)
    anadir_barra_lateral(slide, prs)
    anadir_titulo(slide, "5. EDA: Distribución por tipologías")

    anadir_imagen(slide, "05_tipologia_porcentual.png",
                  left=0.5, top=1.5, height=5.0)

    txBox = slide.shapes.add_textbox(Inches(11.5), Inches(1.8),
                                       Inches(1.7), Inches(4))
    tf = txBox.text_frame
    tf.word_wrap = True
    tf.text = "Residencial:\n66,7%\n\nAparcament:\n28,1%\n\nComercial:\n3,3%\n\nOficina:\n1,4%\n\nTurístic:\n0,3%\n\nEquipaments:\n0,1%"
    for p in tf.paragraphs:
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = COLOR_TEXTO
        p.alignment = PP_ALIGN.CENTER


def slide_08_serpavi(prs):
    slide = nuevo_slide(prs)
    anadir_barra_lateral(slide, prs)
    anadir_titulo(slide, "6. Análisis SERPAVI vs transacciones")

    anadir_imagen(slide, "20_serpavi_vs_transacciones.png",
                  left=0.5, top=1.5, height=5.0)

    txBox = slide.shapes.add_textbox(Inches(11.5), Inches(1.8),
                                       Inches(1.7), Inches(4))
    tf = txBox.text_frame
    tf.word_wrap = True
    tf.text = "Pearson:\nr = -0,07\n(p=0,85)\n\nSpearman:\nρ = -0,27\n(p=0,45)\n\nNO hay\ncorrelación\nsignificativa"
    for p in tf.paragraphs:
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = COLOR_TEXTO
        p.alignment = PP_ALIGN.CENTER


def slide_09_extranjeros(prs):
    slide = nuevo_slide(prs)
    anadir_barra_lateral(slide, prs)
    anadir_titulo(slide, "7. Análisis extranjeros por distrito")

    anadir_imagen(slide, "24_compradores_extranjeros.png",
                  left=0.5, top=1.5, height=5.0)

    txBox = slide.shapes.add_textbox(Inches(11.5), Inches(1.8),
                                       Inches(1.7), Inches(4))
    tf = txBox.text_frame
    tf.word_wrap = True
    tf.text = "Ciutat Vella:\n42,6%\n\nSant Martí:\n30,9%\n\nEixample:\n27,8%\n\nMedia BCN:\n23,0%\n\nSarrià:\n11,5%\n\nLes Corts:\n12,9%"
    for p in tf.paragraphs:
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = COLOR_TEXTO
        p.alignment = PP_ALIGN.CENTER


def slide_10_modelo(prs):
    slide = nuevo_slide(prs)
    anadir_barra_lateral(slide, prs)
    anadir_titulo(slide, "8. Modelo final: XGBoost")

    anadir_caja_destacada(
        slide,
        "XGBoost con objetivo Poisson · 32 features · R² = 0,617",
        top=1.5, height=0.9, color_fondo=COLOR_PRIMARIO,
    )

    items = [
        "Comparativa de modelos",
        "  - Dummy (media): R² = -0,08 · RMSE = 26,53",
        "  - Dummy (media barrio): R² = 0,43 · RMSE = 19,23",
        "  - Ridge: R² = 0,615 · RMSE = 15,82",
        "  - Lasso: R² = 0,605 · RMSE = 16,02",
        "  - XGBoost: R² = 0,617 · RMSE = 15,77",
        "",
        "Mejora sobre baseline: -0,3% RMSE · +0,002 R²",
        "",
        "Justificación de la elección:",
        "  - Overfitting controlado (delta train-test R² = 0,022)",
        "  - Interpretabilidad superior vía SHAP",
        "  - Captura interacciones no lineales",
        "",
        "Interpretación: la mejora es marginal porque las relaciones",
        "son mayormente lineales. Se documenta honestamente.",
    ]
    anadir_bullets(slide, items, top=2.6, tamano=12)


def slide_11_shap_top(prs):
    slide = nuevo_slide(prs)
    anadir_barra_lateral(slide, prs)
    anadir_titulo(slide, "9. SHAP: top features")

    anadir_imagen(slide, "16_shap_bar.png",
                  left=0.5, top=1.5, height=5.0)

    txBox = slide.shapes.add_textbox(Inches(11.5), Inches(1.8),
                                       Inches(1.7), Inches(4))
    tf = txBox.text_frame
    tf.word_wrap = True
    tf.text = "Top 3:\n\n1. rolling_12m\n(0,278)\n\n2. barrio_te\n(0,156)\n\n3. rolling_6m\n(0,097)\n\nEl modelo\nes un\nextrapolador\nde tendencia"
    for p in tf.paragraphs:
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = COLOR_TEXTO
        p.alignment = PP_ALIGN.CENTER


def slide_12_shap_waterfall(prs):
    slide = nuevo_slide(prs)
    anadir_barra_lateral(slide, prs)
    anadir_titulo(slide, "10. SHAP: caso waterfall")

    anadir_imagen(slide, "18_shap_waterfall.png",
                  left=0.5, top=1.5, height=5.0)

    txBox = slide.shapes.add_textbox(Inches(11.5), Inches(1.8),
                                       Inches(1.7), Inches(4))
    tf = txBox.text_frame
    tf.word_wrap = True
    tf.text = "Marina del\nPrat Vermell\nJulio 2025\n\nReal:\n296\n\nPred:\n15,5\n\nResiduo:\n+280,5\n\nEl modelo\nNO captura\npicos\nextremos"
    for p in tf.paragraphs:
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = COLOR_TEXTO
        p.alignment = PP_ALIGN.CENTER


def slide_13_residuos(prs):
    slide = nuevo_slide(prs)
    anadir_barra_lateral(slide, prs)
    anadir_titulo(slide, "11. Análisis de residuos")

    anadir_imagen(slide, "12_analisis_residuos.png",
                  left=0.5, top=1.5, height=5.0)

    txBox = slide.shapes.add_textbox(Inches(11.5), Inches(1.8),
                                       Inches(1.7), Inches(4))
    tf = txBox.text_frame
    tf.word_wrap = True
    tf.text = "Media:\n+1,89\n\nStd:\n15,66\n\nAsimetría:\n4,22\n\nKurtosis:\n52,04\n\nResiduos\nsesgados:\nel modelo\nsobreestima\nbajos y\nsubestima\naltos"
    for p in tf.paragraphs:
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = COLOR_TEXTO
        p.alignment = PP_ALIGN.CENTER


def slide_14_limitaciones(prs):
    slide = nuevo_slide(prs)
    anadir_barra_lateral(slide, prs)
    anadir_titulo(slide, "12. Limitaciones del modelo")

    items = [
        "Del dataset",
        "  - Granularidad temporal mensual (no intra-mensual)",
        "  - Granularidad espacial por barrio (no intra-barrio)",
        "  - Datos del Notariado con 2-3 meses de retraso",
        "  - Falta de precio €/m² en el dataset principal",
        "",
        "Del modelo",
        "  - Subestima valores altos (MAE 67,5 en el rango 100+)",
        "  - Sobreestima valores bajos (en el rango 0-5 sobreestima 8,56)",
        "  - Residuos sesgados (skewness 4,22)",
        "  - No captura picos extremos (caso Marina del Prat Vermell)",
        "  - R² limitado a 0,617",
        "",
        "Del análisis complementario",
        "  - SERPAVI excluye alquileres de personas jurídicas",
        "  - Correlación no implica causalidad",
        "  - Datos de extranjeros parcialmente estimados",
    ]
    anadir_bullets(slide, items, top=1.8, tamano=12)


def slide_15_recomendaciones_inversores(prs):
    slide = nuevo_slide(prs)
    anadir_barra_lateral(slide, prs)
    anadir_titulo(slide, "13. Recomendaciones para inversores")

    items = [
        "1. Mirar el tamaño del distrito, NO el precio del alquiler",
        "   - El análisis SERPAVI vs transacciones muestra correlación nula",
        "   - Factores clave: número de barrios, stock de vivienda",
        "",
        "2. Priorizar barrios con alto rolling_12m (tendencia sostenida)",
        "   - Eixample, Sant Martí, Sant Gervasi",
        "",
        "3. Vigilar barrios emergentes con picos recientes",
        "   - Marina del Prat Vermell: 296 transacciones en julio 2025",
        "   - Diagonal Mar, Poblenou",
        "",
        "4. Evitar barrios con >80% de meses sin actividad",
        "   - Can Peguera, Vallbona, Torre Baró (mercado ilíquido)",
        "",
        "5. Usar el modelo para prever demanda mensual por barrio",
    ]
    anadir_bullets(slide, items, top=1.8, tamano=12)


def slide_16_recomendaciones_admin(prs):
    slide = nuevo_slide(prs)
    anadir_barra_lateral(slide, prs)
    anadir_titulo(slide, "14. Recomendaciones para administraciones")

    items = [
        "1. Vigilar la presión de alquiler en Ciutat Vella",
        "   - 63,7% población extranjera · 44,9% inquilinos extranjeros est.",
        "   - Alquiler más alto (15,55 €/m²) · Riesgo de gentrificación",
        "",
        "2. Monitorizar barrios emergentes",
        "   - Marina del Prat Vermell (desarrollo urbanístico)",
        "   - Diagonal Mar (nueva construcción)",
        "",
        "3. Planificar vivienda protegida en distritos con alta presión",
        "   - Ciutat Vella, Eixample (4.425 transacciones en 2024)",
        "   - Sant Martí (3.291 transacciones en 2024)",
        "",
        "4. Usar el modelo para anticipar tensiones",
        "   - Barrios con rolling_12m creciente y pct_turistico creciente",
        "   - Monitorizar trimestralmente",
        "",
        "5. Publicar datos abiertos a nivel de barrio para investigadores",
    ]
    anadir_bullets(slide, items, top=1.8, tamano=12)


def slide_17_cierre(prs):
    slide = nuevo_slide(prs)
    anadir_barra_lateral(slide, prs)

    # Caja grande de color
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.25), Inches(1.5),
        prs.slide_width - Inches(0.25), Inches(3.5)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = COLOR_PRIMARIO
    shape.line.fill.background()

    txBox = slide.shapes.add_textbox(Inches(1), Inches(2.0),
                                       Inches(11), Inches(2.5))
    tf = txBox.text_frame
    tf.word_wrap = True
    tf.text = "Gracias por su atención"
    p = tf.paragraphs[0]
    p.font.name = FUENTE_TITULO
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = COLOR_BLANCO
    p.alignment = PP_ALIGN.CENTER

    p2 = tf.add_paragraph()
    p2.text = "\nProyecto completo disponible en GitHub:\nreto5-precios-vivienda-cataluna"
    p2.font.size = Pt(16)
    p2.font.color.rgb = COLOR_BLANCO
    p2.alignment = PP_ALIGN.CENTER

    # Pie
    txBox2 = slide.shapes.add_textbox(Inches(1), Inches(5.5),
                                        Inches(11), Inches(1.5))
    tf2 = txBox2.text_frame
    tf2.text = "Trabajo futuro: integración Idescat + SERPAVI a nivel barrio + modelo hurdle"
    p3 = tf2.paragraphs[0]
    p3.font.size = Pt(14)
    p3.font.italic = True
    p3.font.color.rgb = COLOR_TEXTO
    p3.alignment = PP_ALIGN.CENTER


# ============================================================
# MAIN
# ============================================================
def main():
    print("=" * 60)
    print("GENERANDO PRESENTACION 2: RESULTADOS")
    print("=" * 60)

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    print("\nGenerando 17 slides...")
    slide_01_portada(prs)
    slide_02_indice(prs)
    slide_03_resumen_ejecutivo(prs)
    slide_04_ciclo(prs)
    slide_05_estacionalidad(prs)
    slide_06_ranking(prs)
    slide_07_tipologias(prs)
    slide_08_serpavi(prs)
    slide_09_extranjeros(prs)
    slide_10_modelo(prs)
    slide_11_shap_top(prs)
    slide_12_shap_waterfall(prs)
    slide_13_residuos(prs)
    slide_14_limitaciones(prs)
    slide_15_recomendaciones_inversores(prs)
    slide_16_recomendaciones_admin(prs)
    slide_17_cierre(prs)

    SALIDA_PPTX.parent.mkdir(parents=True, exist_ok=True)
    prs.save(SALIDA_PPTX)
    print(f"\nPresentacion guardada: {SALIDA_PPTX}")
    print(f"Total slides: {len(prs.slides)}")


if __name__ == "__main__":
    main()